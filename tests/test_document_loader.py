import pytest
from docx import Document

from app.document_loader import load_document


@pytest.fixture
def word_table_document(tmp_path):
    path = tmp_path / "synthetic-support.docx"
    document = Document()
    document.add_heading("Synthetic support guide", level=1)
    document.add_paragraph("Synthetic policy for loader tests only.")
    document.add_heading("Response targets", level=2)
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [("Priority", "Target"), ("P1", "15 minutes"), ("P2", "1 business hour")],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("Collect diagnostics before escalation.")
    contacts = document.add_table(rows=1, cols=2)
    contacts.cell(0, 0).text = "Network support"
    contacts.cell(0, 1).text = "Diagnostic bundle\nAffected hostnames"
    document.add_paragraph("End of synthetic guide.")
    document.save(path)
    return path


@pytest.fixture
def html_table_document(tmp_path):
    path = tmp_path / "synthetic-support.html"
    path.write_text(
        """<!doctype html>
<html><body>
<nav>Navigation outside evidence</nav>
<main>
<h1>Synthetic support guide</h1>
<p>Synthetic policy for loader tests only.</p>
<h2>Response targets</h2>
<table>
  <tr><th>Priority</th><th>Target</th></tr>
  <tr><td>P1</td><td>15 minutes</td></tr>
  <tr><td>P2</td><td><strong>1</strong> business hour</td></tr>
</table>
<p>Collect diagnostics before escalation.</p>
<table>
  <tr><td>Network support</td><td>Diagnostic bundle<br>Affected hostnames</td></tr>
</table>
<p>End of synthetic guide.</p>
<script>Script outside evidence</script>
</main>
<footer>Footer outside evidence</footer>
</body></html>""",
        encoding="utf-8",
    )
    return path


def test_word_tables_preserve_rows_and_surrounding_order(word_table_document):
    text = load_document(word_table_document)
    expected_blocks = [
        "# Synthetic support guide",
        "Synthetic policy for loader tests only.",
        "## Response targets",
        "Priority | Target\nP1 | 15 minutes\nP2 | 1 business hour",
        "Collect diagnostics before escalation.",
        "Network support | Diagnostic bundle / Affected hostnames",
        "End of synthetic guide.",
    ]
    positions = [text.index(block) for block in expected_blocks]
    assert positions == sorted(positions)


def test_html_tables_preserve_content_order(html_table_document):
    text = load_document(html_table_document)
    expected_content = [
        "Synthetic support guide",
        "Synthetic policy for loader tests only.",
        "Response targets",
        "Priority",
        "Target",
        "P1",
        "15 minutes",
        "P2",
        "1 business hour",
        "Collect diagnostics before escalation.",
        "Network support",
        "Diagnostic bundle",
        "Affected hostnames",
        "End of synthetic guide.",
    ]
    positions = [text.index(content) for content in expected_content]
    assert positions == sorted(positions)
    assert "outside evidence" not in text


def test_html_tables_preserve_explicit_row_relationships(html_table_document):
    text = load_document(html_table_document)
    assert "Priority | Target\nP1 | 15 minutes\nP2 | 1 business hour" in text
    assert "Network support | Diagnostic bundle / Affected hostnames" in text
