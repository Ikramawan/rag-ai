"""Create additive synthetic fixtures; never overwrite existing documents."""

import argparse
import os
import tempfile
from html import escape
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.document_loader import load_document

NOTICE = "Synthetic training documentation only. This is fictional and is not company policy."


def spec(filename, title, sections, table=None):
    return {"filename": filename, "title": title, "sections": sections, "table": table}


SPECS = [
    spec(
        "gcp-labels.md",
        "Synthetic GCP labels",
        [
            (
                "Required labels",
                [
                    "Every training project requires owner, environment and cost-centre labels. The owner label identifies the responsible team, not an individual password or email address.",
                    "The environment label is one of development, test or production. Obtain the cost-centre value from the billing administrator.",
                ],
            ),
            (
                "Corrections",
                [
                    "For a missing or incorrect label, record the project ID and proposed value for platform-team review. This guide does not authorise project creation or deletion."
                ],
            ),
        ],
    ),
    spec(
        "budget-alerts.md",
        "Synthetic GCP budget alerts",
        [
            (
                "Before requesting an alert",
                [
                    "Identify the approved billing account, project ID, monthly threshold and budget owner. The billing administrator reviews budget-alert requests."
                ],
            ),
            (
                "Interpretation",
                [
                    "Budget alerts notify the budget owner when spending reaches a threshold. Alerts do not automatically stop spending or disable paid services. Record the alert time and reported spend before contacting the billing administrator."
                ],
            ),
        ],
    ),
    spec(
        "api-enablement.md",
        "Synthetic API enablement requests",
        [
            (
                "Prerequisites",
                [
                    "The existing project must have platform-team approval and an active billing account. Record the API name, project ID, business purpose and workload owner."
                ],
            ),
            (
                "Procedure",
                [
                    "1. Submit the API enablement request to the platform team.",
                    "2. Wait for the platform team to review the requested API and required permissions.",
                    "3. Record the approval reference before scheduling enablement.",
                    "This guide provides no enablement command or networking procedure.",
                ],
            ),
        ],
    ),
    spec(
        "quota-request.md",
        "Synthetic GCP quota requests",
        [
            (
                "Request evidence",
                [
                    "Record the project ID, region, service, current quota, requested quota and expected workload growth. Include the exact quota error and the business justification."
                ],
            ),
            (
                "Review",
                [
                    "Send the evidence to the platform team. A quota increase requires review; an error does not authorise creating a replacement project. Billing suspension is handled by the billing administrator, not through a quota request."
                ],
            ),
        ],
    ),
    spec(
        "ssh-key-onboarding.md",
        "Synthetic SSH public-key onboarding",
        [
            (
                "Prerequisites",
                [
                    "Obtain approval from the repository owner and use an approved managed device. Prepare the public key and its fingerprint."
                ],
            ),
            (
                "Submission",
                [
                    "Submit only the public key and fingerprint through the access request. Never include a private key, passphrase or recovery code. The identity support team provisions approved access."
                ],
            ),
            (
                "Limits",
                [
                    "This guide contains no key-generation command, server hostname or SSH connection command."
                ],
            ),
        ],
    ),
    spec(
        "git-access.md",
        "Synthetic repository access",
        [
            (
                "New access",
                [
                    "Obtain repository-owner approval. Provide the repository name, requested role, business reason and access expiry date. The identity support team provisions the approved role."
                ],
            ),
            (
                "Troubleshooting",
                [
                    "If repository access is denied after provisioning, record the repository name, account name and exact error. Contact the repository owner to confirm the requested role. Do not share access tokens in a ticket."
                ],
            ),
        ],
    ),
    spec(
        "release-evidence.md",
        "Synthetic release evidence checklist",
        [
            (
                "Prerequisites",
                [
                    "This is an evidence-gathering procedure for a fictional application called Harbor. A change reference, release owner and approved maintenance window are required. It does not execute a deployment."
                ],
            ),
            (
                "Ordered evidence collection",
                [
                    f"{number}. {instruction} Record the observation against the change reference. Preserve the original timestamp and distinguish a verified result from an assumption. If the evidence is unavailable, mark it missing and ask the release owner to resolve the gap before approval."
                    for number, instruction in enumerate(
                        [
                            "Confirm the release owner and change reference.",
                            "Record the intended application version.",
                            "Record the target environment.",
                            "Check that the maintenance window was approved.",
                            "Collect the build identifier from the release record.",
                            "Collect the reviewed test-results reference.",
                            "Identify the business acceptance reviewer.",
                            "Record dependencies listed by the application owner.",
                            "Confirm that the support contact is documented.",
                            "Record the approved communication recipients.",
                            "Collect the documented rollback decision criteria.",
                            "Record the backup-evidence reference if one is required.",
                            "Check that outstanding incident references are listed.",
                            "Identify which observations need follow-up.",
                            "Submit the completed evidence packet for review.",
                            "Wait for explicit approval; this checklist does not deploy the release.",
                        ],
                        start=1,
                    )
                ],
            ),
            (
                "Boundary",
                [
                    "No deployment commands, production credentials or automatic infrastructure actions are provided."
                ],
            ),
        ],
    ),
    spec(
        "rollback-review.md",
        "Synthetic rollback review",
        [
            (
                "Decision request",
                [
                    "For Harbor, the release owner decides whether to request rollback approval. Provide the change reference, observed symptoms, incident reference and documented rollback criteria."
                ],
            ),
            (
                "Evidence",
                [
                    "Record the last known healthy version and the affected environment. Submit the evidence to the release owner. Do not assume rollback approval from a failed test. This guide contains no rollback execution steps or commands."
                ],
            ),
        ],
    ),
    spec(
        "certificate-expiry.md",
        "Synthetic certificate expiry triage",
        [
            (
                "Evidence collection",
                [
                    "Record the affected service, certificate expiry timestamp and exact browser error. Compare the device time with the approved time source before treating the error as a certificate expiry."
                ],
            ),
            (
                "Escalation",
                [
                    "If the certificate is expired, contact the application owner. If the device clock is incorrect, contact the service desk. Never bypass certificate warnings or disable certificate verification."
                ],
            ),
        ],
    ),
    spec(
        "secret-exposure.md",
        "Synthetic exposed-secret reporting",
        [
            (
                "Immediate reporting",
                [
                    "Report a suspected exposed credential to the security response team. Supply the affected system name, discovery time and location of the exposure. Never paste the credential into a ticket or chat."
                ],
            ),
            (
                "Boundary",
                [
                    "The security response team coordinates containment and rotation. This document does not supply credential-rotation commands or authorise the assistant to change credentials."
                ],
            ),
        ],
    ),
    spec(
        "database-read-access.md",
        "Synthetic database read access",
        [
            (
                "Prerequisites",
                [
                    "Obtain approval from the data owner. Record the dataset, business purpose, requested duration and managed-device identity. Only the reporting-reader role is covered here."
                ],
            ),
            (
                "Procedure",
                [
                    "1. Submit the approved access request to the data platform team.",
                    "2. Wait for provisioning confirmation.",
                    "3. Verify that the approval names the intended dataset and expiry date.",
                    "Write access, database hostnames and query commands are not documented.",
                ],
            ),
        ],
    ),
    spec(
        "data-export.md",
        "Synthetic data export review",
        [
            (
                "Approval evidence",
                [
                    "Identify the data owner, approved recipient, dataset classification and intended retention period. The data owner must approve the export request before processing."
                ],
            ),
            (
                "Review packet",
                [
                    "Record the purpose and the proposed destination category without attaching live data. Send the packet to the data governance team. This guide does not define an export command or permission to send company data to an external service."
                ],
            ),
        ],
    ),
    spec(
        "mfa-clock.txt",
        "Synthetic MFA clock mismatch",
        [
            (
                "Symptom",
                [
                    "When a time-based MFA code is rejected and the device time is incorrect, contact the service desk with the device time, timezone and exact error. Do not submit MFA codes or recovery codes.",
                    "An account-lock message belongs to identity support. A clock mismatch does not establish that the account is locked.",
                ],
            )
        ],
    ),
    spec(
        "status-notifications.txt",
        "Synthetic status notifications",
        [
            (
                "Communication",
                [
                    "Harbor status updates are issued by the incident communications owner. Include the incident reference, affected service, current impact and next update time.",
                    "An update time is not a restoration promise. This guide supplies no status-page URL or subscription endpoint.",
                ],
            )
        ],
    ),
    spec(
        "maintenance-timezones.txt",
        "Synthetic maintenance timestamps",
        [
            (
                "Recording windows",
                [
                    "Record each Harbor maintenance window with its calendar date, start time, end time and named timezone. Use Europe/London for the fictional UK support calendar.",
                    "Do not convert UK time to a fixed UTC offset without checking the date. This guide does not change the existing service desk hours or incident response targets.",
                ],
            )
        ],
    ),
    spec(
        "harbor-roles.html",
        "Synthetic Harbor role matrix",
        [
            ("Scope", ["Roles below apply only to the fictional Harbor application."]),
            (
                "Request",
                [
                    "Obtain application-owner approval before requesting a role. Identity support provisions approved access."
                ],
            ),
        ],
        [
            ["Role", "Allowed activity", "Approver"],
            ["Viewer", "Read release records", "Application owner"],
            ["Reviewer", "Review evidence packets", "Release owner"],
            ["Coordinator", "Schedule approved windows", "Application owner"],
        ],
    ),
    spec(
        "backup-retention.html",
        "Synthetic Harbor backup retention",
        [
            (
                "Scope",
                [
                    "These fictional retention values apply only to Harbor backup review fixtures."
                ],
            ),
            (
                "Restore requests",
                [
                    "The data platform team reviews restore requests. A retained backup does not establish that a restore was tested."
                ],
            ),
        ],
        [
            ["Backup class", "Retention", "Owner"],
            ["Daily snapshot", "14 days", "Data platform team"],
            ["Monthly archive", "90 days", "Data owner"],
        ],
    ),
    spec(
        "certificate-routing.html",
        "Synthetic certificate routing matrix",
        [
            (
                "Scope",
                [
                    "Record the service, timestamp and exact certificate error before escalation."
                ],
            ),
            (
                "Boundary",
                [
                    "Do not bypass TLS warnings. No certificate renewal procedure is documented."
                ],
            ),
        ],
        [
            ["Observed condition", "Contact", "Evidence"],
            ["Expired service certificate", "Application owner", "Expiry timestamp"],
            ["Incorrect device time", "Service desk", "Device time and timezone"],
            ["Unknown issuer", "Security response team", "Exact validation error"],
        ],
    ),
    spec(
        "incident-handoff.docx",
        "Synthetic incident handoff",
        [
            (
                "Before handoff",
                [
                    "Prepare the incident reference, impact summary and timestamped observations. State which checks were completed and which remain unverified."
                ],
            ),
            (
                "After the evidence table",
                [
                    "The receiving incident coordinator confirms ownership. An acknowledgement is not a resolution commitment."
                ],
            ),
        ],
        [
            ["Field", "Required evidence"],
            ["Impact", "Affected service and user group"],
            ["Timeline", "Timestamped observations"],
            ["Ownership", "Current coordinator and next contact"],
        ],
    ),
    spec(
        "storage-access.docx",
        "Synthetic storage access matrix",
        [
            (
                "Prerequisites",
                [
                    "Obtain data-owner approval and identify the training storage collection. This guide documents role selection, not an access command."
                ],
            ),
            (
                "Review",
                [
                    "The data platform team provisions the approved role and expiry date. Administrative access is outside this guide."
                ],
            ),
        ],
        [
            ["Role", "Permission", "Duration"],
            ["Report reader", "Read approved reports", "30 days"],
            ["Archive reviewer", "Read archived summaries", "7 days"],
            ["Upload contributor", "Add reviewed training files", "14 days"],
        ],
    ),
    spec(
        "deployment-prerequisites.pdf",
        "Synthetic Harbor deployment prerequisites",
        [
            (
                "Prerequisites",
                [
                    "An approved change reference, release owner, maintenance window and reviewed rollback criteria are required before a Harbor deployment is scheduled."
                ],
            ),
            (
                "Evidence review",
                [
                    "Collect the build identifier and reviewed test-results reference. The release owner reviews missing evidence before authorising scheduling."
                ],
            ),
            (
                "Boundary",
                [
                    "This guide documents prerequisites only. It does not provide deployment commands, production hostnames or permission for the assistant to perform a deployment."
                ],
            ),
        ],
    ),
    spec(
        "diagnostic-redaction.pdf",
        "Synthetic diagnostic redaction",
        [
            (
                "Before preparing a bundle",
                [
                    "Work on a copy of the diagnostic bundle. Retain the original securely for the support owner. This fictional review procedure covers evidence preparation only."
                ],
            ),
            (
                "Review sequence",
                [
                    "First identify the affected service and incident reference. Record the collection time and the device operating system.",
                    "Next review screenshots for passwords, access tokens, MFA recovery codes and unrelated personal information. Remove these items from the review copy.",
                    "Then preserve the exact error text, relevant timestamps and application version. Redaction must not remove the observations needed to distinguish the reported symptom.",
                ],
            ),
            (
                "Handoff",
                [
                    "Finally send the reviewed evidence through the approved support channel. Ask the support owner to confirm receipt. No upload URL or external recipient is documented."
                ],
            ),
        ],
    ),
]


def _write_word(path, item):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = (
        section.right_margin
    ) = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    # compact_reference_guide preset, with a simple left-aligned reference masthead.
    for name, size, before, after, color in [
        ("Normal", 11, 0, 6, "000000"),
        ("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"),
        ("Heading 3", 12, 10, 5, "1F4D78"),
    ]:
        style = document.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before, style.paragraph_format.space_after = (
            Pt(before),
            Pt(after),
        )
        style.paragraph_format.line_spacing = 1.25
    document.add_heading(item["title"], level=1)
    document.add_paragraph(NOTICE)
    for position, (heading, paragraphs) in enumerate(item["sections"]):
        document.add_heading(heading, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if position == 0 and item["table"]:
            rows = item["table"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = [9360 // len(rows[0])] * len(rows[0])
            widths[-1] += 9360 - sum(widths)
            table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:w"), "9360")
            table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "dxa")
            indent = OxmlElement("w:tblInd")
            indent.set(qn("w:w"), "120")
            indent.set(qn("w:type"), "dxa")
            table._tbl.tblPr.append(indent)
            margins = OxmlElement("w:tblCellMar")
            for side, value in [
                ("top", 80),
                ("bottom", 80),
                ("start", 120),
                ("end", 120),
            ]:
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            table._tbl.tblPr.append(margins)
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Inches(width / 1440)
            for row_number, values in enumerate(rows):
                for cell, value, width in zip(
                    table.rows[row_number].cells, values, widths, strict=True
                ):
                    cell.width = Inches(width / 1440)
                    cell.text = value
                    if row_number == 0:
                        fill = OxmlElement("w:shd")
                        fill.set(qn("w:fill"), "E8EEF5")
                        cell._tc.get_or_add_tcPr().append(fill)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
    document.core_properties.author = "Synthetic corpus generator"
    document.save(path)


def _write_pdf(path, item):
    styles = getSampleStyleSheet()
    story = [
        Paragraph(escape(item["title"]), styles["Heading1"]),
        Paragraph(NOTICE, styles["BodyText"]),
    ]
    for heading, paragraphs in item["sections"]:
        if item["filename"] == "diagnostic-redaction.pdf" and heading == "Handoff":
            story.append(PageBreak())
        story.extend([Spacer(1, 12), Paragraph(escape(heading), styles["Heading2"])])
        for paragraph in paragraphs:
            story.extend(
                [Paragraph(escape(paragraph), styles["BodyText"]), Spacer(1, 6)]
            )
    SimpleDocTemplate(
        str(path),
        pagesize=(612, 792),
        leftMargin=72,
        rightMargin=72,
        topMargin=72,
        bottomMargin=72,
        title=item["title"],
        author="Synthetic corpus generator",
    ).build(story)


def _write_document(path, item):
    if path.suffix == ".docx":
        _write_word(path, item)
    elif path.suffix == ".pdf":
        _write_pdf(path, item)
    elif path.suffix == ".html":
        blocks = [
            "<!doctype html><html><body><main>",
            f"<h1>{escape(item['title'])}</h1><p>{NOTICE}</p>",
        ]
        for position, (heading, paragraphs) in enumerate(item["sections"]):
            blocks.append(f"<h2>{escape(heading)}</h2>")
            blocks.extend(f"<p>{escape(paragraph)}</p>" for paragraph in paragraphs)
            if position == 0 and item["table"]:
                blocks.append("<table>")
                for row_number, row in enumerate(item["table"]):
                    tag = "th" if row_number == 0 else "td"
                    blocks.append(
                        "<tr>"
                        + "".join(f"<{tag}>{escape(value)}</{tag}>" for value in row)
                        + "</tr>"
                    )
                blocks.append("</table>")
        blocks.append("</main></body></html>")
        path.write_text("\n".join(blocks), encoding="utf-8")
    else:
        blocks = ["# " + item["title"], NOTICE]
        for heading, paragraphs in item["sections"]:
            blocks.extend(["## " + heading, *paragraphs])
        path.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")


def generate(output: Path) -> list[Path]:
    output = Path(output)
    targets = [output / item["filename"] for item in SPECS]
    if any(path.exists() or path.is_symlink() for path in targets):
        raise FileExistsError(
            "A corpus filename already exists; nothing was overwritten."
        )
    output.mkdir(parents=True, exist_ok=True)
    created = []
    with tempfile.TemporaryDirectory(
        prefix="synthetic-corpus-", dir=output.parent
    ) as staging:
        for item in SPECS:
            path = Path(staging) / item["filename"]
            _write_document(path, item)
            if NOTICE not in load_document(path):
                raise ValueError(
                    f"Synthetic notice missing after extraction: {path.name}"
                )
        try:
            for target in targets:
                os.link(Path(staging) / target.name, target)
                created.append(target)
        except OSError:
            for target in created:
                target.unlink()
            raise
    return created


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    for path in generate(args.output):
        print(path)


if __name__ == "__main__":
    main()
